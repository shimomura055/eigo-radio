# -*- coding: utf-8 -*-
r"""KPMG 集計とExcel出力(API 0)。Trial専用。KP6M(er098)の集計・Excel処理を踏襲。
入力: runs/, gen/, eval/t1, eval/t2, validation.json, paid_calls.jsonl, 前回 er098_output/kp_six_model_trial_01/{validation,summary}.json
出力: summary.json, t2_mechanical.json, docs/research/FAMILY-X-KP-SELECTION-PROMPT-MIN-AND-GENERALIZATION-TRIAL-02/KP_PROMPT_MIN_GEN_COMPARISON_02.xlsx"""
import json, re, sys, statistics
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
import er101_kp_prompt_min_gen_trial_02_common as K
import er098_kp_six_model_trial_01_validate as V
C = K.C
OUT = K.OUT
PREV = C.ROOT / "er098_output" / "kp_six_model_trial_01"
XLSX = C.ROOT / "docs" / "research" / "FAMILY-X-KP-SELECTION-PROMPT-MIN-AND-GENERALIZATION-TRIAL-02" / "KP_PROMPT_MIN_GEN_COMPARISON_02.xlsx"
GREEN = PatternFill("solid", fgColor="C6EFCE"); YELLOW = PatternFill("solid", fgColor="FFEB9C"); RED = PatternFill("solid", fgColor="FFC7CE")
GREY = PatternFill("solid", fgColor="D9D9D9"); HEAD = PatternFill("solid", fgColor="DDEBF7"); BLUE = PatternFill("solid", fgColor="BDD7EE")
BOLD = Font(bold=True); WRAP = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="999999"); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
FAMILY = {"n1_gpt-6-luna": "openai", "n2_gemini-3.1-flash-lite": "gemini", "n3_gemini-3.5-flash-lite": "gemini", "n4_claude-haiku-5-5": "anthropic", "n5_claude-sonnet-5-5": "anthropic"}
EVFAM = {"sol": "openai", "sonnet": "anthropic"}
SRC = list(K.MODELS.keys())
def lab(s): return "%d %s" % (K.MODELS[s]["no"], K.MODELS[s]["model_id"])
def avg(xs):
    xs = [x for x in xs if x is not None]; return sum(xs) / len(xs) if xs else None
def jl(p): return json.loads(p.read_text(encoding="utf8"))
EXAMPLES = [("think up what happens next", r"think\w*\s+up"), ("formal agreements with their creators", r"agreement"), ("alone caused the different results", r"caus")]
PLACEH = {"someone", "something", "somebody", "someone's", "somebody's", "something's", "one's", "oneself"}

def fill_new(v):
    if v["strict_pass"]: return GREEN
    if v["match_type"] == "NG_INFLECTION_ONLY" and v["within_5"]: return YELLOW
    return RED
def fill_prev(v):
    if v["match_type"] in ("A_EXACT", "B_GEN_UNIQUE") and v["within_5"] and v["duplicate_of"] is None: return GREEN
    if v["match_type"] in ("B_GEN_AMBIGUOUS", "NG_INFLECTION_ONLY") and v["within_5"]: return YELLOW
    return RED

def t2_mech(before, after, reason):
    bt = before.split(); at = after.split()
    r = dict(before=before, after=after, words_after=len(at), within_5=0 < len(at) <= 5, changed=(before.strip() != after.strip()))
    bs = [V.stem(t) for t in bt]; aset = []
    kinds = []
    if r["changed"]:
        bl = [t.lower() for t in bt]; al = [t.lower() for t in at]
        if any(t.lower().strip(",.") in PLACEH for t in at) and not any(t.lower() in PLACEH for t in bt): kinds.append("someone/something等へ置換")
        if [V.stem(t) for t in al] == bs and al != bl: kinds.append("原形化/単複/活用のみ")
        elif len(al) < len(bl) and all(t in bl for t in al): kinds.append("語の削除")
        elif not kinds or True:
            if not (kinds and kinds[0] == "原形化/単複/活用のみ"): kinds.append("語の置換/追加を含む" if any(V.stem(t) not in bs and t.strip(",.") not in PLACEH for t in at) else "語順/その他")
    r["change_kinds"] = kinds
    added = [t for t in at if V.stem(t) not in bs and t.lower().strip(",.") not in PLACEH]
    r["added_tokens"] = added
    r["added_flag"] = bool(added)
    return r

def main():
    val = jl(OUT / "validation.json"); pval = jl(PREV / "validation.json"); psum = jl(PREV / "summary.json")
    man = jl(OUT / "inputs" / "manifest_from_kp6m.json"); units = [u for u in man["units"] if u.get("copy_path")]
    runs = {}
    for rp in (OUT / "runs").glob("*/*.json"):
        d = jl(rp); runs[(d["model_key"], d["unit_id"])] = d
    gens = {u["unit_id"]: jl(OUT / "gen" / ("%s.json" % u["unit_id"])) for u in units}
    bmap = jl(OUT / "eval" / "blind_map.json")
    ev1 = {}; ev2 = {}
    for ev in EVFAM:
        for u in units:
            p1 = OUT / "eval" / "t1" / ev / ("%s.json" % u["unit_id"]); p2 = OUT / "eval" / "t2" / ev / ("%s.json" % u["unit_id"])
            if p1.exists(): ev1[(ev, u["unit_id"])] = jl(p1)
            if p2.exists(): ev2[(ev, u["unit_id"])] = jl(p2)
    vset = {(s["source"], s["unit_id"]): s for s in val["sets"]}
    pvset = {(s["source"], s["unit_id"]): s for s in pval["sets"]}
    # ---- T1 scores
    cs = defaultdict(dict); pad = defaultdict(dict)
    for (ev, uid), d in ev1.items():
        p = d.get("parsed")
        if not p: continue
        for st in p.get("sets", []):
            src = bmap[uid].get(st.get("set_id"))
            if not src: continue
            pad[(src, uid)][ev] = bool(st.get("padding"))
            for c in st.get("candidates", []):
                try: sc = {k: int(c[k]) for k in "VENCUSX"}
                except Exception: continue
                sc["old3"] = (sc["V"] + sc["E"] + sc["N"]) / 6.0; sc["min4"] = (sc["C"] + sc["U"] + sc["S"] + sc["X"]) / 8.0; sc["note"] = c.get("note", "")
                cs[(src, uid, int(c["idx"]))][ev] = sc
    rows = []
    for (src, uid), s in vset.items():
        for v in s["candidates"]:
            per = cs.get((src, uid, v["rank"]), {})
            main = [e for e in per if EVFAM[e] != FAMILY[src]]; ref = [e for e in per if EVFAM[e] == FAMILY[src]]
            rows.append(dict(source=src, unit_id=uid, rank=v["rank"], v=v, per=per, main=main, ref=ref))
    summary = {"models": {}, "notes": {}}
    for src in SRC:
        rs = [r for r in rows if r["source"] == src]; sets = [s for k, s in vset.items() if k[0] == src]
        n = len(rs); vs = [r["v"] for r in rs]
        m = dict(sets=len(sets), tech_fail_sets=sum(s["status"] != "OK" for s in sets), candidates=n, within_5=sum(v["within_5"] for v in vs), over5=sum(not v["within_5"] for v in vs),
                 exact=sum(v["match_type"] == "A_EXACT" for v in vs), infl=sum(v["match_type"] == "NG_INFLECTION_ONLY" for v in vs),
                 ng_other=sum(v["match_type"] in ("NG_NO_MATCH", "NG_ONLY_PLACEHOLDER", "NG_EMPTY") for v in vs), placeholder_dev=sum(v["match_type"] == "NG_PLACEHOLDER_DEVIATION" for v in vs),
                 dup=sum(v["duplicate_of"] is not None for v in vs), strict=sum(v["strict_pass"] for v in vs), self_count_reported=sum(v["self_count"] is not None for v in vs),
                 mismatch=sum(v["self_count_mismatch"] for v in vs), sets_lt5=sum(1 for s in sets if len(s["candidates"]) < 5), sets_gt5=sum(1 for s in sets if len(s["candidates"]) > 5),
                 sets_all_strict=sum(1 for s in sets if s["candidates"] and len(s["candidates"]) <= 5 and all(c["strict_pass"] for c in s["candidates"])),
                 draft_over5_mention_sets=sum(1 for s in sets if s["parse_info"].get("draft_over5_mention")), avg_words=avg([v["words_actual"] for v in vs]),
                 words_dist={str(k): sum(v["words_actual"] == k for v in vs) for k in range(1, 7)})
        for k in "VENCUSX":
            m["main_" + k] = avg([r["per"][e][k] for r in rs for e in r["main"]])
            m["ref_" + k] = avg([r["per"][e][k] for r in rs for e in r["ref"]])
        m["main_old3"] = avg([r["per"][e]["old3"] for r in rs for e in r["main"]]); m["main_min4"] = avg([r["per"][e]["min4"] for r in rs for e in r["main"]])
        m["ref_old3"] = avg([r["per"][e]["old3"] for r in rs for e in r["ref"]]); m["ref_min4"] = avg([r["per"][e]["min4"] for r in rs for e in r["ref"]])
        m["eval_n_main"] = sum(len(r["main"]) for r in rs)
        m["padding_rate"] = avg([1.0 if p[e] else 0.0 for (s2, u2), p in pad.items() if s2 == src for e in p if EVFAM[e] != FAMILY[src]])
        # 前回(KP6M)同モデル
        pk = K.MODELS[src]["prev_key"]; pm = psum["models"][pk]
        m["prev"] = dict(candidates=pm["candidates"], strict=pm["strict"], within_5=pm["within_5"], main_V=pm["main_V"], main_E=pm["main_E"], main_N=pm["main_N"],
                         main_old3=(pm["main_V"] + pm["main_E"] + pm["main_N"]) / 6.0, avg_words=avg([c["words_actual"] for k, s in pvset.items() if k[0] == pk for c in s["candidates"]]),
                         mismatch=pm["mismatch"], sets_lt5=pm["sets_lt5"], infl=pm["infl_only"], ng_other=pm["ng_other"], generalized=pm["generalized_confirmed"], cost_jpy=pm["cost"]["cost_jpy"])
        cost = dict(calls=0, in_tokens=0, out_tokens=0, reasoning_tokens=0, cost_jpy=0.0)
        for (mk, uid), d in runs.items():
            if mk == src and d.get("ok"):
                u = d["usage"]; cost["calls"] += 1; cost["in_tokens"] += u.get("input_tokens") or 0; cost["out_tokens"] += u.get("output_tokens") or 0; cost["reasoning_tokens"] += u.get("reasoning_tokens") or 0; cost["cost_jpy"] += d["cost_jpy"]
        m["cost"] = cost; m["cost_per_set"] = cost["cost_jpy"] / 7
        m["by_article"] = {a: dict(n=sum(1 for r in rs if r["unit_id"].startswith(a + "_")), strict=sum(1 for r in rs if r["unit_id"].startswith(a + "_") and r["v"]["strict_pass"]),
                                   main_min4=avg([r["per"][e]["min4"] for r in rs if r["unit_id"].startswith(a + "_") for e in r["main"]])) for a in "ABCDEFG"}
        m["violations"] = ["%s:%s(%s)" % (r["unit_id"], r["v"]["phrase"], r["v"]["match_type"] if r["v"]["within_5"] else "OVER5") for r in rs if not r["v"]["strict_pass"]]
        summary["models"][src] = m
    # 例示3件の追跡
    tr = []
    for ex, pat in EXAMPLES:
        for (pk_, uid), s in pvset.items():
            for c in s["candidates"]:
                if ex.lower() == c["phrase"].lower():
                    t = dict(example=ex, prev_model=pk_, unit_id=uid, prev_words=c["words_actual"], new=[])
                    for src in SRC:
                        if K.MODELS[src]["prev_key"] != pk_: continue
                        ns = vset.get((src, uid))
                        if ns: t["new_model"] = src; t["new"] = [dict(phrase=x["phrase"], words=x["words_actual"], match=x["match_type"]) for x in ns["candidates"] if re.search(pat, x["phrase"], re.I)]
                        t["new_all"] = [x["phrase"] for x in ns["candidates"]] if ns else []
                    tr.append(t)
    summary["example_tracking"] = tr
    # ---- T2 mechanical + eval
    t2 = []
    for u in units:
        g = gens[u["unit_id"]]; sets_l = vset[("n1_gpt-6-luna", u["unit_id"])]["candidates"]
        p = g["parsed"] or []
        ctxs = re.findall(r"^(\d+)\. (.+)\n   本文中の該当文\(参考情報\): (.+)$", g["user_message"], flags=re.M)
        for i, (c, (no, bef, sent)) in enumerate(zip(sets_l, ctxs)):
            item = p[i] if i < len(p) else {}
            mech = t2_mech(c["phrase"], item.get("after", ""), item.get("reason", ""))
            mech.update(unit_id=u["unit_id"], rank=c["rank"], sentence=sent, reason=item.get("reason", ""), before_echo_ok=(item.get("before", "").strip() == c["phrase"].strip()), count_ok=(len(p) == len(sets_l)), no_ok=(item.get("no") == c["rank"]))
            ev = {}
            for e in EVFAM:
                d = ev2.get((e, u["unit_id"]))
                if d and d.get("parsed"):
                    its = {int(x["no"]): x for x in d["parsed"].get("items", []) if "no" in x}
                    if c["rank"] in its: ev[e] = its[c["rank"]]
            mech["eval"] = ev
            t2.append(mech)
    (OUT / "t2_mechanical.json").write_text(json.dumps(t2, ensure_ascii=False, indent=1), encoding="utf8")
    def b(x): return bool(x)
    t2s = dict(n=len(t2), changed=sum(r["changed"] for r in t2), within_5=sum(r["within_5"] for r in t2), before_echo_ok=sum(r["before_echo_ok"] for r in t2), added_flag=sum(r["added_flag"] for r in t2),
               reselection_or_count_ng=sum((not r["before_echo_ok"]) or (not r["count_ok"]) or (not r["no_ok"]) for r in t2), by_eval={})
    for e in EVFAM:
        ri = [r["eval"][e] for r in t2 if e in r["eval"]]
        t2s["by_eval"][e] = dict(n=len(ri), needs_change=sum(b(x.get("needs_change")) for x in ri), changed_ok_pass=sum(x.get("changed_ok") == "PASS" for x in ri), keep_pass=sum(x.get("keep_point") == "PASS" for x in ri),
                                 natural_pass=sum(x.get("natural") == "PASS" for x in ri), meaning_change=sum(b(x.get("meaning_change")) for x in ri), over_change=sum(b(x.get("over_change")) for x in ri),
                                 missed=sum(b(x.get("missed")) for x in ri), added_words=sum(b(x.get("added_words")) for x in ri))
    both = [r for r in t2 if len(r["eval"]) == 2]
    def both_cnt(k, val_=True): return sum(1 for r in both if all((x.get(k) in (val_,) if not isinstance(val_, str) else x.get(k) == val_) for x in r["eval"].values()))
    t2s["both"] = dict(n=len(both), changed_ok_pass_both=sum(1 for r in both if all(x.get("changed_ok") == "PASS" for x in r["eval"].values())), keep_pass_both=sum(1 for r in both if all(x.get("keep_point") == "PASS" for x in r["eval"].values())),
                       natural_pass_both=sum(1 for r in both if all(x.get("natural") == "PASS" for x in r["eval"].values())),
                       all4_pass_both=sum(1 for r in both if all(x.get("changed_ok") == "PASS" and x.get("keep_point") == "PASS" and x.get("natural") == "PASS" for x in r["eval"].values()) and r["within_5"]),
                       missed_any=sum(1 for r in both if any(x.get("missed") for x in r["eval"].values())), missed_both=sum(1 for r in both if all(x.get("missed") for x in r["eval"].values())),
                       over_any=sum(1 for r in both if any(x.get("over_change") for x in r["eval"].values())), over_both=sum(1 for r in both if all(x.get("over_change") for x in r["eval"].values())),
                       meaning_any=sum(1 for r in both if any(x.get("meaning_change") for x in r["eval"].values())), meaning_both=sum(1 for r in both if all(x.get("meaning_change") for x in r["eval"].values())),
                       needs_both=sum(1 for r in both if all(x.get("needs_change") for x in r["eval"].values())), needs_any=sum(1 for r in both if any(x.get("needs_change") for x in r["eval"].values())),
                       changed_and_needs_any=sum(1 for r in both if r["changed"] and any(x.get("needs_change") for x in r["eval"].values())),
                       unchanged_and_needs_both=sum(1 for r in both if (not r["changed"]) and all(x.get("needs_change") for x in r["eval"].values())),
                       changed_and_not_needed_both=sum(1 for r in both if r["changed"] and all(not x.get("needs_change") for x in r["eval"].values())),
                       needs_agree=sum(1 for r in both if len({bool(x.get("needs_change")) for x in r["eval"].values()}) == 1))
    summary["t2"] = t2s
    # ---- 費用
    gen_cost = dict(calls=0, in_tokens=0, out_tokens=0, cost_jpy=0.0)
    for d in gens.values():
        if d.get("ok"): gen_cost["calls"] += 1; gen_cost["in_tokens"] += d["usage"].get("input_tokens") or 0; gen_cost["out_tokens"] += d["usage"].get("output_tokens") or 0; gen_cost["cost_jpy"] += d["cost_jpy"]
    summary["gen_cost"] = gen_cost
    evc = {}
    for ph, evs in (("t1", ev1), ("t2", ev2)):
        for (e, uid), d in evs.items():
            x = evc.setdefault(ph + ":" + e, dict(calls=0, in_tokens=0, out_tokens=0, reasoning_tokens=0, cost_jpy=0.0, parse_errors=0, model_returned=set()))
            if d.get("ok"):
                u = d["usage"]; x["calls"] += 1; x["in_tokens"] += u.get("input_tokens") or 0; x["out_tokens"] += u.get("output_tokens") or 0; x["reasoning_tokens"] += u.get("reasoning_tokens") or 0
                x["cost_jpy"] += d["cost_jpy"]; x["model_returned"].add(d.get("model_returned")); x["parse_errors"] += 0 if d.get("parsed") else 1
    for x in evc.values(): x["model_returned"] = sorted(y for y in x["model_returned"] if y)
    summary["eval_cost"] = evc
    sel_total = sum(summary["models"][s]["cost"]["cost_jpy"] for s in SRC); ev_total = sum(x["cost_jpy"] for x in evc.values())
    summary["totals"] = dict(selection=sel_total, generalization=gen_cost["cost_jpy"], evaluation=ev_total, sum=sel_total + gen_cost["cost_jpy"] + ev_total, ledger=C.ledger_total_jpy())
    luna_sel = summary["models"]["n1_gpt-6-luna"]["cost"]["cost_jpy"]
    summary["production_unit_cost"] = dict(luna_new_prompt_selection_per_set=luna_sel / 7, haiku_gen_per_set=gen_cost["cost_jpy"] / 7, two_step_per_set=(luna_sel + gen_cost["cost_jpy"]) / 7,
                                           production_existing_per_set=sum((c.get("selector") or {}).get("cost_jpy_recorded", (c.get("selector") or {}).get("cost_jpy_recomputed", 0)) + (c.get("canonicalization") or {}).get("cost_jpy_recomputed", 0)
                                                                          for c in [u.get("production_cost", {}) for u in units]) / 7)
    pp = psum["models"]["production"]; summary["production_prev"] = dict(strict=pp["strict"], candidates=pp["candidates"], main_V=pp["main_V"], main_E=pp["main_E"], main_N=pp["main_N"], main_G=pp["main_G"], old3=(pp["main_V"] + pp["main_E"] + pp["main_N"]) / 6)
    (OUT / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=1, default=str), encoding="utf8")
    build_xlsx(units, summary, vset, pvset, rows, t2, runs, gens, evc)
    print(json.dumps(dict(totals=summary["totals"], unit=summary["production_unit_cost"]), ensure_ascii=False, indent=1))
    for s in SRC:
        m = summary["models"][s]
        print(s, "n=%d strict=%d w5=%d avgw=%.2f(prev %.2f) mismatch=%d lt5=%d old3=%.3f(prev %.3f) min4=%.3f C/U/S/X=%.2f/%.2f/%.2f/%.2f cost=%.3f" % (m["candidates"], m["strict"], m["within_5"], m["avg_words"], m["prev"]["avg_words"], m["mismatch"], m["sets_lt5"],
              m["main_old3"] or -1, m["prev"]["main_old3"], m["main_min4"] or -1, m["main_C"] or -1, m["main_U"] or -1, m["main_S"] or -1, m["main_X"] or -1, m["cost"]["cost_jpy"]))
    print(json.dumps(t2s, ensure_ascii=False))

def build_xlsx(units, summary, vset, pvset, rows, t2, runs, gens, evc):
    wb = Workbook()
    # ---- Sheet1 T1比較
    ws = wb.active; ws.title = "T1 モデル比較"
    ws.append(["Trial① 最小化Prompt x 5モデル(セル=『表現 (語数)』)。各モデルの直下に『前回KP6M(旧Prompt)の同モデル結果』を灰色文字行で並記。色=機械検証(緑=厳格合格[本文一致・5語以内・重複なし]/黄=活用形のみ相違/赤=NG)。Production採用未承認のTrial"]); ws["A1"].font = BOLD
    ws.append([]); ws.append(["記事/レベル", "選定元", "KP1", "KP2", "KP3", "KP4", "KP5", "備考"])
    for c in ws[3]: c.font = BOLD; c.fill = HEAD; c.border = BORDER
    gfont = Font(color="777777", italic=True)
    for u in units:
        uid = u["unit_id"]; ws.append(["記事%s %s / %s (%s)" % (u["article_id"], u["article_name"], u["level"], u["level_internal"])])
        for c in ws[ws.max_row]: c.fill = BLUE
        ws.cell(row=ws.max_row, column=1).font = BOLD
        for src in SRC:
            r = ws.max_row + 1; ws.cell(row=r, column=1, value=u["level"]); ws.cell(row=r, column=2, value="【今回】" + lab(src)).font = BOLD
            s = vset.get((src, uid))
            for k, v in enumerate(s["candidates"][:5]):
                cell = ws.cell(row=r, column=3 + k, value="%s (%d)" % (v["phrase"], v["words_actual"])); cell.fill = fill_new(v); cell.alignment = WRAP
                cell.comment = Comment("判定=%s; 5語以内=%s; 原文span=%s; 自己申告語数=%s" % (v["match_type"], v["within_5"], v.get("span_text"), v.get("self_count")), "er101")
            note = "件数=%d" % len(s["candidates"]) + (" (5件未満)" if len(s["candidates"]) < 5 else "")
            if s["parse_info"].get("manual_final_list") or s["parse_info"].get("phrase_overrides") or s["parse_info"].get("override_group") is not None: note += " / 抽出を手動補正(TRIAL_RESULT_02 8節)"
            if s["parse_info"].get("draft_over5_mention"): note += " / 出力中に6語以上の草稿・除外記述あり"
            ws.cell(row=r, column=8, value=note)
            pk = K.MODELS[src]["prev_key"]; ps = pvset.get((pk, uid)); r2 = ws.max_row + 1
            ws.cell(row=r2, column=1, value=u["level"]); c2 = ws.cell(row=r2, column=2, value="(前回KP6M・旧Prompt)" + lab(src)); c2.font = gfont
            if ps:
                for k, v in enumerate(ps["candidates"][:5]):
                    cell = ws.cell(row=r2, column=3 + k, value="%s (%d)" % (v["phrase"], v["words_actual"])); cell.fill = fill_prev(v); cell.font = gfont; cell.alignment = WRAP
                ws.cell(row=r2, column=8, value="前回: 旧Prompt(一般化条項あり)。語数平均 %.2f" % avg([c["words_actual"] for c in ps["candidates"]])).font = gfont
    for col, w in zip("ABCDEFGH", (34, 40, 30, 30, 30, 30, 30, 44)): ws.column_dimensions[col].width = w
    ws.freeze_panes = "C4"
    # ---- Sheet2 T1品質評価
    w2 = wb.create_sheet("T1 品質評価")
    w2.append(["Trial① モデル別 指示遵守・最小化品質(14本文。分母=出力候補数)。評価尺度の注意: V/E/Nは前回と同一定義だが別バッチ評価(同一尺度の厳密比較ではなく参考)。C/U/S/Xは今回新設(前回なし)。前回の総合平均はG(一般化)を含むため今回と直接比較しない"]); w2["A1"].font = BOLD
    hdr = ["選定元", "セット数(技術失敗)", "候補数", "5語以内 率", "本文厳密一致 率(A_EXACT)", "活用形のみ相違(NG)", "NG(言い換え/創作)", "Prompt逸脱(placeholder)", "5語超", "重複", "厳格合格 率(指示遵守率)", "全件厳格のセット数", "5件未満セット",
           "自己申告語数あり", "自己申告不一致", "出力に6語以上の草稿記述があるセット", "平均語数(今回)", "平均語数(前回KP6M同モデル)", "語数分布1/2/3/4/5/6語(今回)",
           "主 V", "主 E", "主 N", "V/E/N平均(0-1・前回と同定義)", "前回 V/E/N平均(同定義・別バッチ)", "主 C 明確さ", "主 U 不要語なし", "主 S 最短形", "主 X 不自然さなし", "最小化4観点平均(0-1)", "参考(同系統)最小化4観点平均", "padding率(主)", "選定実費JPY(14本文)", "1記事セット単価JPY"]
    w2.append(hdr)
    for c in w2[2]: c.font = BOLD; c.fill = HEAD; c.alignment = WRAP; c.border = BORDER
    for src in SRC:
        m = summary["models"][src]; n = m["candidates"]
        w2.append([lab(src), "%d(%d)" % (m["sets"], m["tech_fail_sets"]), n, m["within_5"] / n, m["exact"] / n, m["infl"], m["ng_other"], m["placeholder_dev"], m["over5"], m["dup"], m["strict"] / n, "%d/%d" % (m["sets_all_strict"], m["sets"]), m["sets_lt5"],
                   m["self_count_reported"], m["mismatch"], m["draft_over5_mention_sets"], m["avg_words"], m["prev"]["avg_words"], "/".join(str(m["words_dist"][str(k)]) for k in range(1, 7)),
                   m["main_V"], m["main_E"], m["main_N"], m["main_old3"], m["prev"]["main_old3"], m["main_C"], m["main_U"], m["main_S"], m["main_X"], m["main_min4"], m["ref_min4"], m["padding_rate"], m["cost"]["cost_jpy"], m["cost_per_set"]])
        rr = w2.max_row
        for col in (4, 5, 11, 31): w2.cell(row=rr, column=col).number_format = "0.0%"
        for col in (17, 18, 20, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30): w2.cell(row=rr, column=col).number_format = "0.00"
        for col in (32, 33): w2.cell(row=rr, column=col).number_format = "0.000"
    pp = summary["production_prev"]
    w2.append([]); w2.append(["参考: 前回KP6Mの既存Production行(DB hybrid方式・別方式)", "", pp["candidates"], "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", pp["main_V"], pp["main_E"], pp["main_N"], pp["old3"]])
    w2.append([]); w2.append(["定義: 主評価=選定元と異系統の評価者(Luna=Sonnet / Haiku・Sonnet=Sol / Gemini=Sol・Sonnet平均)。評価者=gpt-6.1-sol, claude-sonnet-5-5(前回と同一)。V/E/N=前回rubric(各0-2)、C/U/S/X=今回追加(各0-2)。厳格合格=5語以内かつ本文厳密一致(placeholder・活用形差はNG)かつ重複なし。"])
    w2.append([]); w2.append(["記事別 厳格合格数/候補数 と 最小化4観点平均(主)"]); w2.cell(row=w2.max_row, column=1).font = BOLD
    w2.append(["選定元"] + ["記事%s" % a for a in "ABCDEFG"] + [""] + ["記事%s 最小化平均" % a for a in "ABCDEFG"])
    for c in w2[w2.max_row]: c.font = BOLD; c.fill = HEAD
    for src in SRC:
        bb = summary["models"][src]["by_article"]; w2.append([lab(src)] + ["%d/%d" % (bb[a]["strict"], bb[a]["n"]) for a in "ABCDEFG"] + [""] + [bb[a]["main_min4"] for a in "ABCDEFG"])
        for col in range(10, 17): w2.cell(row=w2.max_row, column=col).number_format = "0.00"
    w2.append([]); w2.append(["違反内容(厳格不合格): unit:候補(判定)"]); w2.cell(row=w2.max_row, column=1).font = BOLD
    for src in SRC: w2.append([lab(src), "; ".join(summary["models"][src]["violations"]) or "なし"]); w2.cell(row=w2.max_row, column=2).alignment = WRAP
    w2.append([]); w2.append(["ユーザー例示3件の追跡(前回の長い表現が今回どうなったか。同モデル・同記事の今回選定と比較)"]); w2.cell(row=w2.max_row, column=1).font = BOLD
    w2.append(["例示", "前回モデル", "unit", "前回語数", "今回の同モデル同記事での該当候補(語数)", "今回の同モデル同記事の全選定"])
    for c in w2[w2.max_row]: c.font = BOLD; c.fill = HEAD
    for t in summary["example_tracking"]:
        w2.append([t["example"], t["prev_model"], t["unit_id"], t["prev_words"], "; ".join("%s(%d,%s)" % (x["phrase"], x["words"], x["match"]) for x in t["new"]) or "(該当する語句なし=選ばれず)", " / ".join(t.get("new_all", []))])
    w2.column_dimensions["A"].width = 34
    for i in range(2, 34): w2.column_dimensions[get_column_letter(i)].width = 15
    w2.freeze_panes = "B3"
    # ---- Sheet3 T2 Before/After
    w3 = wb.create_sheet("T2 一般化Before-After")
    w3.append(["Trial② Luna選定(Trial①出力そのまま=Before)→claude-haiku-5-5一般化(After)。再選定なし。評価=gpt-6.1-sol / claude-sonnet-5-5(両者とも一般化担当と異系統)。変更件数の多さは加点しない。色: 緑=2評価者とも全PASS、赤=いずれかNG/要確認、青=変更なし"]); w3["A1"].font = BOLD
    w3.append(["記事/レベル", "#", "Before(Luna選定)", "After(Haiku一般化)", "変更内容(機械分類)", "一般化担当の変更理由", "5語以内(機械)", "追加語(機械)", "要変更?(Sol/Sonnet)", "一般化の適切さ(Sol/Sonnet)", "学習ポイント維持(Sol/Sonnet)", "英語の自然さ(Sol/Sonnet)", "適用漏れ(Sol/Sonnet)", "過剰変更(Sol/Sonnet)", "意味変化(Sol/Sonnet)", "評価者メモ"])
    for c in w3[2]: c.font = BOLD; c.fill = HEAD; c.alignment = WRAP; c.border = BORDER
    uname = {u["unit_id"]: "記事%s %s / %s" % (u["article_id"], u["article_name"], u["level"]) for u in units}
    def pair(r, k, f=lambda x: x):
        return "/".join(str(f(r["eval"][e].get(k))) if e in r["eval"] else "-" for e in ("sol", "sonnet"))
    for r in t2:
        ev = r["eval"]
        allpass = len(ev) == 2 and all(x.get("changed_ok") == "PASS" and x.get("keep_point") == "PASS" and x.get("natural") == "PASS" for x in ev.values()) and r["within_5"] and not r["added_flag"]
        notes = " | ".join("%s:%s" % (e, x.get("note", "")) for e, x in ev.items() if x.get("note"))
        w3.append([uname[r["unit_id"]], r["rank"], r["before"], r["after"], ("変更なし" if not r["changed"] else "; ".join(r["change_kinds"])), r["reason"], "OK" if r["within_5"] else "5語超", ", ".join(r["added_tokens"]) if r["added_flag"] else "",
                   pair(r, "needs_change", lambda x: "要" if x else "不要"), pair(r, "changed_ok"), pair(r, "keep_point"), pair(r, "natural"), pair(r, "missed", lambda x: "漏れ" if x else "-"), pair(r, "over_change", lambda x: "過剰" if x else "-"),
                   pair(r, "meaning_change", lambda x: "変化" if x else "-"), notes])
        rr = w3.max_row
        fl = GREEN if allpass else (RED if (len(ev) < 2 or any(x.get("changed_ok") == "NG" or x.get("keep_point") == "NG" or x.get("natural") == "NG" for x in ev.values()) or not r["within_5"]) else YELLOW)
        if not r["changed"] and allpass: fl = PatternFill("solid", fgColor="BDD7EE")
        for col in (3, 4): w3.cell(row=rr, column=col).fill = fl
        for col in (4, 6, 16): w3.cell(row=rr, column=col).alignment = WRAP
    w3.append([]); t = summary["t2"]; bth = t["both"]
    w3.append(["集計(70組)", "", "変更あり %d / 変更なし %d" % (t["changed"], t["n"] - t["changed"]), "5語以内 %d/%d" % (t["within_5"], t["n"]), "機械の追加語検出 %d" % t["added_flag"], "再選定・件数不一致等 %d" % t["reselection_or_count_ng"]])
    w3.append(["", "", "2評価者とも一般化の適切さPASS %d/%d" % (bth["changed_ok_pass_both"], bth["n"]), "学習ポイント維持 両PASS %d" % bth["keep_pass_both"], "自然さ 両PASS %d" % bth["natural_pass_both"], "3項目全PASS+5語以内 %d" % bth["all4_pass_both"]])
    w3.append(["", "", "適用漏れ(いずれか/両者) %d/%d" % (bth["missed_any"], bth["missed_both"]), "過剰変更(いずれか/両者) %d/%d" % (bth["over_any"], bth["over_both"]), "意味変化(いずれか/両者) %d/%d" % (bth["meaning_any"], bth["meaning_both"]), "要変更の判定一致 %d/%d" % (bth["needs_agree"], bth["n"])])
    for col, w in zip("ABCDEFGHIJKLMNOP", (30, 4, 26, 28, 24, 50, 9, 12, 12, 12, 12, 12, 12, 12, 12, 50)): w3.column_dimensions[col].width = w
    w3.freeze_panes = "C3"
    # ---- Sheet4 T1検証詳細
    w4 = wb.create_sheet("T1 検証詳細")
    w4.append(["選定元", "記事", "レベル", "順位", "候補(出力)", "語数(実測)", "自己申告語数", "自己申告不一致", "5語以内", "本文一致種別", "原文span", "文字位置", "重複", "厳格合格", "主 V", "主 E", "主 N", "主 C", "主 U", "主 S", "主 X", "主 V/E/N平均", "主 最小化4観点平均", "参考(同系統)最小化4観点平均", "評価者メモ"])
    for c in w4[1]: c.font = BOLD; c.fill = HEAD; c.alignment = WRAP; c.border = BORDER
    uu = {u["unit_id"]: u for u in units}
    for u in units:
        for src in SRC:
            for r in [x for x in rows if x["source"] == src and x["unit_id"] == u["unit_id"]]:
                v = r["v"]; per = r["per"]
                def mv(k): return avg([per[e][k] for e in r["main"]]) if r["main"] else None
                w4.append([lab(src), "記事%s" % u["article_id"], u["level"], v["rank"], v["phrase"], v["words_actual"], v["self_count"], "不一致" if v["self_count_mismatch"] else "", "OK" if v["within_5"] else "5語超", v["match_type"], v.get("span_text"),
                           ("%s-%s" % (v.get("span_start"), v.get("span_end"))) if v.get("span_start") is not None else "", "重複(#%s)" % v["duplicate_of"] if v["duplicate_of"] else "", "合格" if v["strict_pass"] else "不合格",
                           mv("V"), mv("E"), mv("N"), mv("C"), mv("U"), mv("S"), mv("X"), mv("old3"), mv("min4"), avg([per[e]["min4"] for e in r["ref"]]) if r["ref"] else None, " | ".join("%s:%s" % (e, per[e]["note"]) for e in per if per[e].get("note"))])
                w4.cell(row=w4.max_row, column=10).fill = fill_new(v)
    for i, w in enumerate((26, 8, 10, 6, 34, 8, 8, 10, 8, 24, 36, 14, 8, 8, 7, 7, 7, 7, 7, 7, 7, 10, 10, 12, 30), 1): w4.column_dimensions[get_column_letter(i)].width = w
    w4.freeze_panes = "F2"; w4.auto_filter.ref = w4.dimensions
    # ---- Sheet5 T2検証詳細
    w5 = wb.create_sheet("T2 検証詳細")
    w5.append(["記事/レベル", "#", "Before", "After", "本文中の文脈文(一般化担当・評価者へ提示)", "Before一致(echo)", "件数/番号一致", "語数(After)", "5語以内", "変更あり", "追加語(機械)", "Sol: 要変更/適切/維持/自然/漏れ/過剰/意味変化/追加語/メモ", "Sonnet: 同左"])
    for c in w5[1]: c.font = BOLD; c.fill = HEAD; c.alignment = WRAP; c.border = BORDER
    for r in t2:
        def f(e):
            x = r["eval"].get(e)
            return "" if not x else "%s/%s/%s/%s/%s/%s/%s/%s/%s" % ("要" if x.get("needs_change") else "不要", x.get("changed_ok"), x.get("keep_point"), x.get("natural"), x.get("missed"), x.get("over_change"), x.get("meaning_change"), x.get("added_words"), x.get("note", ""))
        w5.append([uname[r["unit_id"]], r["rank"], r["before"], r["after"], r["sentence"], "OK" if r["before_echo_ok"] else "不一致", "OK" if (r["count_ok"] and r["no_ok"]) else "NG", r["words_after"], "OK" if r["within_5"] else "5語超", "あり" if r["changed"] else "なし",
                   ", ".join(r["added_tokens"]), f("sol"), f("sonnet")])
        for col in (5, 12, 13): w5.cell(row=w5.max_row, column=col).alignment = WRAP
    for i, w in enumerate((30, 4, 26, 28, 60, 10, 10, 8, 8, 8, 14, 60, 60), 1): w5.column_dimensions[get_column_letter(i)].width = w
    w5.freeze_panes = "E2"
    # ---- Sheet6 費用・実行条件
    w6 = wb.create_sheet("費用・実行条件")
    w6.append(["費用・実行条件(実費=APIが返したusage x 公式/登録単価。為替USD/JPY=%.0f はプロジェクト既定[er009:20]、実為替ではない)" % C.USD_JPY]); w6["A1"].font = BOLD
    w6.append(["工程", "model ID(指定→返却)", "呼出数", "入力token計", "出力token計(推論含む)", "うち推論token", "単価 in/cached/out (USD/1M)", "単価出典", "実費JPY", "1記事セット(Std+Adv=2本文)当たりJPY", "備考"])
    for c in w6[2]: c.font = BOLD; c.fill = HEAD; c.alignment = WRAP; c.border = BORDER
    for src in SRC:
        m = summary["models"][src]; c = m["cost"]; p = K.PRICES[K.MODELS[src]["model_id"]]
        ret = sorted({d.get("model_returned") for (mk, _), d in runs.items() if mk == src and d.get("ok")} - {None})
        w6.append(["Trial① 選定 " + K.MODELS[src]["label"], "%s → %s" % (K.MODELS[src]["model_id"], ", ".join(ret)), c["calls"], c["in_tokens"], c["out_tokens"], c["reasoning_tokens"] or "", "%s / %s / %s" % p[:3], p[3], round(c["cost_jpy"], 4), round(c["cost_jpy"] / 7, 4), "14本文=7記事セット"])
    gc = summary["gen_cost"]; p = K.PRICES["claude-haiku-5-5"]
    gret = sorted({d.get("model_returned") for d in gens.values() if d.get("ok")} - {None})
    w6.append(["Trial② 一般化 Claude Haiku 5.5", "claude-haiku-5-5 → %s" % ", ".join(gret), gc["calls"], gc["in_tokens"], gc["out_tokens"], "", "%s / %s / %s" % p[:3], p[3], round(gc["cost_jpy"], 4), round(gc["cost_jpy"] / 7, 4), "1本文=1呼出(5件まとめて)"])
    for k, x in evc.items():
        e = k.split(":")[1]; mid = "gpt-6.1-sol" if e == "sol" else "claude-sonnet-5-5"; p = K.PRICES[mid]
        w6.append(["独立評価 %s %s" % ("Trial①" if k.startswith("t1") else "Trial②", e), "%s → %s" % (mid, ", ".join(x["model_returned"])), x["calls"], x["in_tokens"], x["out_tokens"], x["reasoning_tokens"] or "", "%s / %s / %s" % p[:3], p[3], round(x["cost_jpy"], 4), round(x["cost_jpy"] / 7, 4), "評価費(選定費・一般化費と別)"])
    t = summary["totals"]; pu = summary["production_unit_cost"]
    w6.append([]); w6.append(["Trial実費合計", "", "", "", "", "", "", "", round(t["sum"], 3), "", "選定 %.3f + 一般化 %.3f + 評価 %.3f(台帳paid_calls.jsonl通算 %.3f)" % (t["selection"], t["generalization"], t["evaluation"], t["ledger"])])
    w6.append([]); w6.append(["量産原価(1記事セット=Standard+Advanced。評価費を含まない運用時の費用)"]); w6.cell(row=w6.max_row, column=1).font = BOLD
    w6.append(["今回: Luna新Prompt選定 JPY/セット", round(pu["luna_new_prompt_selection_per_set"], 4)]); w6.append(["今回: Haiku一般化 JPY/セット", round(pu["haiku_gen_per_set"], 4)])
    w6.append(["今回: 2工程合計 JPY/セット", round(pu["two_step_per_set"], 4)]); w6.append(["既存Production KP(DB hybrid selector+canonicalization・実測) JPY/セット", round(pu["production_existing_per_set"], 4), "方式差: 全文渡し(今回) vs DB hybrid shortlist(既存)。直接の優劣比較ではない"])
    w6.append([]); w6.append(["Trial①適用Prompt(全文固定・全5モデル共通・無改変。sha256=%s)" % C.sha_text((OUT / "inputs" / "prompt_fixed.txt").read_text(encoding="utf8"))]); w6.cell(row=w6.max_row, column=1).font = BOLD
    w6.append([(OUT / "inputs" / "prompt_fixed.txt").read_text(encoding="utf8")]); w6.cell(row=w6.max_row, column=1).alignment = WRAP; w6.row_dimensions[w6.max_row].height = 330
    w6.append(["Trial①受け渡し調整(事後報告): 変更前=Prompt本文のみ / 変更後=Prompt本文は1字も変更せず末尾に『\\n\\n---\\n\\n【記事本文】\\n』+本文を後置(KP6Mと同一) / 理由=本文の受け渡し位置が必要 / 影響なし=選定条件・順序・優先順位は同一、出力形式指示の追加なし。"]); w6.cell(row=w6.max_row, column=1).alignment = WRAP; w6.row_dimensions[w6.max_row].height = 60
    gp = (OUT / "inputs" / "gen_prompt_fixed.txt").read_text(encoding="utf8")
    w6.append([]); w6.append(["Trial②適用一般化Prompt(委任文§4-2を一字一句固定。sha256=%s)" % C.sha_text(gp)]); w6.cell(row=w6.max_row, column=1).font = BOLD
    w6.append([gp]); w6.cell(row=w6.max_row, column=1).alignment = WRAP; w6.row_dimensions[w6.max_row].height = 150
    w6.append(["Trial②受け渡し調整(事後報告): 変更前=一般化Prompt本文のみ / 変更後=本文の後ろに『\\n\\n---\\n\\n【Key Phrase一覧】』+番号付きKey Phrase(各項目に『本文中の該当文(参考情報)』を全70件同一形式で付与)+『【出力形式(受け渡し用の指定)】JSON配列のみ、{no,before,after,reason}、入力と同順・同件数』を付与 / 理由=パースと対応追跡のため、文脈文は固有名詞・this/its等の依存判断材料として全件同一形式で付与 / 影響=一般化条件・5語以内・再選定禁止の指示は不変。"]); w6.cell(row=w6.max_row, column=1).alignment = WRAP; w6.row_dimensions[w6.max_row].height = 110
    w6.append([]); w6.append(["使用本文(KP6Mと同一・sha256一致をrun前にassert)"]); w6.cell(row=w6.max_row, column=1).font = BOLD
    w6.append(["unit", "記事", "レベル(内部)", "本文path", "commit", "本文copy sha256", "元article sha256", "文字数", "語数"])
    for c in w6[w6.max_row]: c.font = BOLD; c.fill = HEAD
    for u in units: w6.append([u["unit_id"], u["article_name"], "%s(%s)" % (u["level"], u["level_internal"]), u["copy_path"], u["article_commit"], u["copy_sha256"], u["article_sha256_file"], u["article_chars"], u["article_words"]])
    w6.append([]); w6.append(["単価(USD/1M tokens, input/cached/output[thinking込み])と出典"]); w6.cell(row=w6.max_row, column=1).font = BOLD
    for k, p in K.PRICES.items(): w6.append([k, p[0], p[1], p[2], p[3]])
    w6.append([]); w6.append(["実行条件: 温度・推論設定はprovider既定(未送信)、max出力token=%d(全モデル同一)、system promptなし、各モデル1本文1回(品質理由の再生成なし)、技術的再試行は記録。" % C.MAX_OUT])
    w6.column_dimensions["A"].width = 60
    for col in "BCDEFGHIJK": w6.column_dimensions[col].width = 24
    XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX); print("saved", XLSX)

if __name__ == "__main__":
    sys.exit(main())
