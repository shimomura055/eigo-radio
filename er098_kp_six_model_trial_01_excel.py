# -*- coding: utf-8 -*-
"""FAMILY-X-KP-SELECTION-SIX-MODEL-COMPARISON-TRIAL-01 委任_01: 集計とExcel出力(API 0)。Trial専用。
入力: inputs/manifest.json, runs/, validation.json, eval/, paid_calls.jsonl
出力: summary.json, docs/research/FAMILY-X-KP-SELECTION-SIX-MODEL-COMPARISON-TRIAL-01/KP_MODEL_COMPARISON_01.xlsx"""
import json, sys
from pathlib import Path
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
import er098_kp_six_model_trial_01_common as C

OUT = C.OUT
XLSX = C.ROOT / "docs" / "research" / "FAMILY-X-KP-SELECTION-SIX-MODEL-COMPARISON-TRIAL-01" / "KP_MODEL_COMPARISON_01.xlsx"
GREEN = PatternFill("solid", fgColor="C6EFCE"); YELLOW = PatternFill("solid", fgColor="FFEB9C"); RED = PatternFill("solid", fgColor="FFC7CE")
GREY = PatternFill("solid", fgColor="D9D9D9"); HEAD = PatternFill("solid", fgColor="DDEBF7"); BLUE = PatternFill("solid", fgColor="BDD7EE")
BOLD = Font(bold=True); WRAP = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="999999"); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CONF = {"A_EXACT", "B_GEN_UNIQUE"}
FAMILY = {"m1_gpt-6-luna": "openai", "m2_gpt-6.1-sol": "openai", "production": "openai", "m3_gemini-3.1-flash-lite": "gemini", "m4_gemini-3.5-flash-lite": "gemini",
          "m5_claude-haiku-5-5": "anthropic", "m6_claude-sonnet-5-5": "anthropic"}
EVFAM = {"sol": "openai", "sonnet": "anthropic"}
SRC_ORDER = ["production"] + list(C.MODELS.keys())
def src_label(s): return "既存Production" if s == "production" else "%d %s" % (C.MODELS[s]["no"], C.MODELS[s]["model_id"])
THRESH = 0.75

def fill_for(v):
    if v["match_type"] in CONF and v["within_5"] and v["duplicate_of"] is None: return GREEN
    if v["match_type"] in ("B_GEN_AMBIGUOUS", "NG_INFLECTION_ONLY") and v["within_5"]: return YELLOW
    return RED

def load():
    man = json.loads((OUT / "inputs" / "manifest.json").read_text(encoding="utf8"))
    val = json.loads((OUT / "validation.json").read_text(encoding="utf8"))
    runs = {}
    for rp in (OUT / "runs").glob("*/*.json"):
        d = json.loads(rp.read_text(encoding="utf8")); runs[(d["model_key"], d["unit_id"])] = d
    evs = {}; bmap = {}
    if (OUT / "eval" / "blind_map.json").exists(): bmap = json.loads((OUT / "eval" / "blind_map.json").read_text(encoding="utf8"))
    for ev in EVFAM:
        for rp in (OUT / "eval" / ev).glob("*.json") if (OUT / "eval" / ev).exists() else []:
            d = json.loads(rp.read_text(encoding="utf8")); evs[(ev, d["unit_id"])] = d
    return man, val, runs, evs, bmap

def eval_scores(evs, bmap):
    """-> {(source, unit, rank): {ev: dict(G,V,E,N,score,needs_gen)}}, padding {(source,unit): {ev: bool}}"""
    cs = defaultdict(dict); pad = defaultdict(dict)
    for (ev, uid), d in evs.items():
        p = d.get("parsed")
        if not p: continue
        for st in p.get("sets", []):
            src = bmap.get(uid, {}).get(st.get("set_id"))
            if not src: continue
            pad[(src, uid)][ev] = bool(st.get("padding"))
            for c in st.get("candidates", []):
                try:
                    g, v, e, n = (int(c[k]) for k in ("G", "V", "E", "N"))
                except Exception: continue
                cs[(src, uid, int(c["idx"]))][ev] = dict(G=g, V=v, E=e, N=n, score=(g + v + e + n) / 8.0, needs_gen=bool(c.get("needs_generalization")), note=c.get("note", ""))
    return cs, pad

def main_ref(src, per):
    """主評価(異系統)と参考評価(同系統)。Gemini行は2系統平均=主評価。"""
    fam = FAMILY[src]; main = [e for e in per if EVFAM[e] != fam]; ref = [e for e in per if EVFAM[e] == fam]
    return main, ref

def avg(xs):
    xs = [x for x in xs if x is not None]; return sum(xs) / len(xs) if xs else None

def main():
    man, val, runs, evs, bmap = load()
    cs, pad = eval_scores(evs, bmap)
    units = [u for u in man["units"] if u.get("copy_path")]
    uorder = [u["unit_id"] for u in units]
    vset = {(s["source"], s["unit_id"]): s for s in val["sets"]}
    # 候補ごとの統合レコード
    rows = []
    for (src, uid), s in vset.items():
        for v in s["candidates"]:
            per = cs.get((src, uid, v["rank"]), {})
            main_evs, ref_evs = main_ref(src, per)
            main_score = avg([per[e]["score"] for e in main_evs]); ref_score = avg([per[e]["score"] for e in ref_evs])
            dims_ok = all(min(per[e][k] for k in "GVEN") >= 1 for e in main_evs) if main_evs else None
            strict = bool(v["strict_pass"])
            comp = bool(strict and main_score is not None and main_score >= THRESH and dims_ok)
            rows.append(dict(source=src, unit_id=uid, rank=v["rank"], phrase=v["phrase"], words=v["words_actual"], self_count=v.get("self_count"), mismatch=v["self_count_mismatch"],
                             within_5=v["within_5"], match_type=v["match_type"], span=v.get("span_text"), start=v.get("span_start"), end=v.get("span_end"),
                             gen=v.get("generalization"), dup=v["duplicate_of"], strict=strict, main_score=main_score, ref_score=ref_score, comp=comp,
                             per=per, main_evs=main_evs, ref_evs=ref_evs, note=v.get("note")))
    by = defaultdict(list)
    for r in rows: by[r["source"]].append(r)
    summary = {"models": {}, "threshold_comprehensive": THRESH}
    for src in SRC_ORDER:
        rs = by.get(src, []); sets = [s for k, s in vset.items() if k[0] == src]
        n = len(rs)
        m = dict(sets=len(sets), sets_ok=sum(1 for s in sets if s["status"] == "OK"), tech_fail_sets=sum(1 for s in sets if s["status"] != "OK"), candidates=n,
                 within_5=sum(r["within_5"] for r in rs), over5=sum(not r["within_5"] for r in rs),
                 confirmed=sum(r["match_type"] in CONF for r in rs), ambiguous=sum(r["match_type"] == "B_GEN_AMBIGUOUS" for r in rs),
                 infl_only=sum(r["match_type"] == "NG_INFLECTION_ONLY" for r in rs), ng_other=sum(r["match_type"] in ("NG_NO_MATCH", "NG_ONLY_PLACEHOLDER", "NG_EMPTY") for r in rs),
                 generalized_confirmed=sum(r["match_type"] == "B_GEN_UNIQUE" for r in rs), mismatch=sum(r["mismatch"] for r in rs), self_count_reported=sum(r["self_count"] is not None for r in rs),
                 dup=sum(r["dup"] is not None for r in rs), strict=sum(r["strict"] for r in rs), comp=sum(r["comp"] for r in rs),
                 sets_lt5=sum(1 for s in sets if s["status"] == "OK" and len(s["candidates"]) < 5), sets_gt5=sum(1 for s in sets if s["status"] == "OK" and len(s["candidates"]) > 5),
                 sets_all_strict=sum(1 for s in sets if s["status"] == "OK" and s["candidates"] and len(s["candidates"]) <= 5 and all(c["strict_pass"] for c in s["candidates"])))
        ms = [r["main_score"] for r in rs if r["main_score"] is not None]
        m["main_avg"] = avg(ms); m["ref_avg"] = avg([r["ref_score"] for r in rs]); m["eval_n"] = len(ms)
        for k in "GVEN":
            vals = []
            for r in rs:
                for e in r["main_evs"]:
                    vals.append(r["per"][e][k])
            m["main_" + k] = avg(vals)
        gvals = []
        for r in rs:
            for e in r["main_evs"]:
                if r["per"][e]["needs_gen"]: gvals.append(r["per"][e]["G"])
        m["main_G_when_needed"] = avg(gvals); m["needs_gen_n"] = len(gvals)
        pads = [pad[(src, s["unit_id"])] for s in sets if (src, s["unit_id"]) in pad]
        pv = []
        for p in pads:
            main_evs, _ = main_ref(src, p); pv += [p[e] for e in main_evs]
        m["padding_rate"] = avg([1.0 if x else 0.0 for x in pv]); m["padding_n"] = len(pv)
        # 記事別
        m["by_article"] = {}
        for aid in "ABCDEFG":
            ra = [r for r in rs if r["unit_id"].startswith(aid + "_")]
            m["by_article"][aid] = dict(n=len(ra), strict=sum(r["strict"] for r in ra), main_avg=avg([r["main_score"] for r in ra]), comp=sum(r["comp"] for r in ra),
                                        violations=[("%s:%s(%s)" % (r["unit_id"], r["phrase"], r["match_type"] if r["within_5"] else "OVER5")) for r in ra if not r["strict"]])
        # 費用
        cost = dict(in_tokens=0, out_tokens=0, reasoning_tokens=0, cost_jpy=0.0, calls=0)
        if src != "production":
            for (mk, uid), d in runs.items():
                if mk == src and d.get("ok"):
                    u = d["usage"]; cost["in_tokens"] += u.get("input_tokens") or 0; cost["out_tokens"] += u.get("output_tokens") or 0; cost["reasoning_tokens"] += u.get("reasoning_tokens") or 0
                    cost["cost_jpy"] += d["cost_jpy"]; cost["calls"] += 1
        m["cost"] = cost
        summary["models"][src] = m
    # 評価費用
    evcost = {}
    for (ev, uid), d in evs.items():
        e = evcost.setdefault(ev, dict(calls=0, in_tokens=0, out_tokens=0, reasoning_tokens=0, cost_jpy=0.0, parse_errors=0, model_returned=set()))
        if d.get("ok"):
            u = d["usage"]; e["calls"] += 1; e["in_tokens"] += u.get("input_tokens") or 0; e["out_tokens"] += u.get("output_tokens") or 0; e["reasoning_tokens"] += u.get("reasoning_tokens") or 0
            e["cost_jpy"] += d["cost_jpy"]; e["model_returned"].add(d.get("model_returned"))
            if not d.get("parsed"): e["parse_errors"] += 1
    for e in evcost.values(): e["model_returned"] = sorted(x for x in e["model_returned"] if x)
    summary["eval_cost"] = evcost
    summary["total_jpy"] = C.ledger_total_jpy()
    # Production既存実費
    pc = []
    for u in units:
        c = u.get("production_cost", {}); sel = c.get("selector"); can = c.get("canonicalization")
        sj = sel.get("cost_jpy_recorded", sel.get("cost_jpy_recomputed")) if isinstance(sel, dict) else None
        cj = can.get("cost_jpy_recomputed") if isinstance(can, dict) else None
        pc.append(dict(unit_id=u["unit_id"], selector_jpy=sj, canonicalization_jpy=cj, kind=sel.get("kind") if isinstance(sel, dict) else "未確認", selector_model=sel.get("model_id") if isinstance(sel, dict) else None))
    summary["production_cost"] = pc
    (OUT / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=1, default=str), encoding="utf8")

    wb = Workbook()
    # ---- シート1 比較表
    ws = wb.active; ws.title = "比較表"
    ws.append(["Family X Key Phrase選定 6モデル比較(同一Prompt・同一本文)。セル=「表現 (語数)」。色=機械検証。セルにマウスを載せると原文span・判定。※Production採用未承認のTrial"])
    ws.append(["凡例:"]); lg = ws.max_row
    for i, (f, t) in enumerate([(GREEN, "本文一致 確認済み かつ 5語以内"), (YELLOW, "要確認(一般化span複数該当)/活用形・所有格代名詞のみ本文と相違"), (RED, "NG(本文に連続して存在しない/創作・言い換え/5語超/重複)"), (GREY, "技術的失敗(API) または 出力なし")]):
        c = ws.cell(row=lg, column=2 + i * 2, value=t); c.fill = f; c.alignment = WRAP
    ws.cell(row=lg, column=1, value="凡例"); ws.row_dimensions[lg].height = 45
    ws.append([]); ws.append(["記事/レベル", "選定元", "KP1", "KP2", "KP3", "KP4", "KP5", "備考(件数・技術的失敗・注記)"])
    for c in ws[ws.max_row]: c.font = BOLD; c.fill = HEAD; c.border = BORDER
    for u in units:
        uid = u["unit_id"]
        ws.append(["記事%s %s / %s (%s)" % (u["article_id"], u["article_name"], u["level"], u["level_internal"])])
        for c in ws[ws.max_row]: c.fill = BLUE
        ws.cell(row=ws.max_row, column=1).font = BOLD
        for src in SRC_ORDER:
            s = vset.get((src, uid))
            r = ws.max_row + 1
            ws.cell(row=r, column=1, value=u["level"]); ws.cell(row=r, column=2, value=src_label(src)).font = BOLD
            note = ""
            if s is None or s["status"] != "OK":
                for k in range(5): ws.cell(row=r, column=3 + k).fill = GREY
                note = "技術的失敗/未実行"
            else:
                for k, v in enumerate(s["candidates"][:8]):
                    col = 3 + k
                    if col > 7:
                        note += " 6件目以降: %s(%d)" % (v["phrase"], v["words_actual"]); continue
                    cell = ws.cell(row=r, column=col, value="%s (%d)" % (v["phrase"], v["words_actual"]))
                    cell.fill = fill_for(v); cell.alignment = WRAP
                    cm = "判定=%s; 5語以内=%s; 原文span=%s; 位置=%s-%s; 自己申告語数=%s" % (v["match_type"], v["within_5"], v.get("span_text"), v.get("span_start"), v.get("span_end"), v.get("self_count"))
                    if v.get("generalization"): cm += "; 一般化=" + ", ".join("%s<-%s" % (g["placeholder"], g["original"]) for g in v["generalization"])
                    cell.comment = Comment(cm, "er098")
                for k in range(len(s["candidates"]), 5): pass
                note = ("件数=%d" % len(s["candidates"])) + (" (5件未満)" if len(s["candidates"]) < 5 else "") + note
                if src != "production":
                    d = runs.get((src, uid))
                    if d and d.get("attempts") and len(d["attempts"]) > 1: note += " / API再試行%d回" % (len(d["attempts"]) - 1)
            ws.cell(row=r, column=8, value=note)
    for col, w in zip("ABCDEFGH", (34, 28, 30, 30, 30, 30, 30, 40)): ws.column_dimensions[col].width = w
    ws.freeze_panes = "C5"
    # ---- シート2 品質評価サマリ
    w2 = wb.create_sheet("品質評価サマリ")
    w2.append(["モデル別 指示遵守率・選定品質(全14本文。分母=出力された候補数)。既存Productionは別方式(shortlist+canonicalization)の参考値"]); w2["A1"].font = BOLD
    hdr = ["選定元", "セット数(技術的失敗)", "候補数", "5語以内 率", "本文一致 確認済み 率", "うち一般化(置換)で確認", "要確認(複数span)", "活用形のみ相違(NG扱い)", "NG(言い換え/創作等)", "5語超", "重複", "自己申告語数あり", "自己申告不一致",
           "5件未満セット", "5件超セット", "全件strict合格セット(/セット)", "厳格版 合格KP数", "総合版 合格KP数", "主評価 平均(0-1)", "参考評価(同系統) 平均", "主 G", "主 V", "主 E", "主 N", "一般化必要箇所の主G平均", "padding率(主)"]
    w2.append(hdr)
    for c in w2[2]: c.font = BOLD; c.fill = HEAD; c.alignment = WRAP; c.border = BORDER
    def pct(a, b): return (a / b) if b else None
    for src in SRC_ORDER:
        m = summary["models"][src]; n = m["candidates"]
        w2.append([src_label(src), "%d(%d)" % (m["sets"], m["tech_fail_sets"]), n, pct(m["within_5"], n), pct(m["confirmed"], n), m["generalized_confirmed"], m["ambiguous"], m["infl_only"], m["ng_other"], m["over5"], m["dup"],
                   m["self_count_reported"], m["mismatch"], m["sets_lt5"], m["sets_gt5"], "%d/%d" % (m["sets_all_strict"], m["sets_ok"]), m["strict"], m["comp"], m["main_avg"], m["ref_avg"],
                   m["main_G"], m["main_V"], m["main_E"], m["main_N"], m["main_G_when_needed"], m["padding_rate"]])
        for col in (4, 5): w2.cell(row=w2.max_row, column=col).number_format = "0.0%"
        w2.cell(row=w2.max_row, column=26).number_format = "0.0%"
        for col in (19, 20, 21, 22, 23, 24, 25): w2.cell(row=w2.max_row, column=col).number_format = "0.00"
    w2.append([]); w2.append(["定義: 厳格版=5語以内かつ本文一致が確認済み(A_EXACT/B_GEN_UNIQUE)かつ重複なし。総合版=厳格版かつ主評価>=%.2f かつ主評価のG/V/E/Nが全て1以上。主評価=選定元と異系統の評価者(Gemini行は2系統平均)。" % THRESH])
    w2.append(["注: 既存Productionは原形化・短縮(canonicalization)済みのため厳密規則で『活用形のみ相違』『言い換え』と判定される行がある。方式差の観測値でありProduction仕様違反ではない。"])
    w2.append([]); w2.append(["記事別 厳格版合格数/候補数 と 主評価平均"]); w2.cell(row=w2.max_row, column=1).font = BOLD
    w2.append(["選定元"] + ["記事%s" % a for a in "ABCDEFG"] + ["", "(主評価平均)"] + ["記事%s" % a for a in "ABCDEFG"])
    for c in w2[w2.max_row]: c.font = BOLD; c.fill = HEAD
    for src in SRC_ORDER:
        b = summary["models"][src]["by_article"]
        w2.append([src_label(src)] + ["%d/%d" % (b[a]["strict"], b[a]["n"]) for a in "ABCDEFG"] + ["", ""] + [b[a]["main_avg"] for a in "ABCDEFG"])
        for col in range(11, 18): w2.cell(row=w2.max_row, column=col).number_format = "0.00"
    w2.append([]); w2.append(["違反内容(厳格版不合格の候補、モデル別): unit:候補(判定)"]); w2.cell(row=w2.max_row, column=1).font = BOLD
    for src in SRC_ORDER:
        vs = [x for a in "ABCDEFG" for x in summary["models"][src]["by_article"][a]["violations"]]
        w2.append([src_label(src), "; ".join(vs) if vs else "なし"]); w2.cell(row=w2.max_row, column=2).alignment = WRAP
    w2.column_dimensions["A"].width = 30
    for i in range(2, 27): w2.column_dimensions[get_column_letter(i)].width = 16
    w2.freeze_panes = "B3"
    # ---- シート3 費用サマリ
    w3 = wb.create_sheet("費用サマリ")
    w3.append(["費用サマリ(実費=APIが返したusage x 公式/登録単価。為替=USD/JPY %.0f はプロジェクト既定値[er009:20]。推定ではなく実測usage)" % C.USD_JPY]); w3["A1"].font = BOLD
    w3.append(["選定元", "model ID(返却)", "呼出数", "入力token計", "出力token計(推論含む)", "うち推論token", "単価 in/cached/out (USD/1M)", "単価出典", "選定実費 JPY(14本文)", "1記事セット単価 JPY(Std+Adv)", "7記事セット合計 JPY",
               "厳格版KP 1件当たり JPY", "総合版KP 1件当たり JPY", "備考"])
    for c in w3[2]: c.font = BOLD; c.fill = HEAD; c.alignment = WRAP; c.border = BORDER
    for src in list(C.MODELS.keys()):
        m = summary["models"][src]; c = m["cost"]; p = C.PRICES[C.MODELS[src]["model_id"]]
        ret = sorted({d.get("model_returned") for (mk, _), d in runs.items() if mk == src and d.get("ok")} - {None})
        w3.append([src_label(src), ", ".join(ret), c["calls"], c["in_tokens"], c["out_tokens"], c["reasoning_tokens"] if c["reasoning_tokens"] else "", "%s / %s / %s" % (p[0], p[1], p[2]), p[3],
                   round(c["cost_jpy"], 4), round(c["cost_jpy"] / 7, 4), round(c["cost_jpy"], 4), (round(c["cost_jpy"] / m["strict"], 4) if m["strict"] else "n/a"),
                   (round(c["cost_jpy"] / m["comp"], 4) if m["comp"] else "n/a"), "評価費用は含まない(下表)"])
    w3.append(["注: 『うち推論token』= APIが分離して返した値(OpenAI reasoning_tokens)。Gemini 3.xは既定でthoughtsTokenCount=0、Claude(Haiku/Sonnet 5.5)は既定でthinking blockが付き出力tokenに含まれるが分離値は返らない(出力token計に包含)。1記事セット=Standard+Advanced=2本文、7記事セット合計=14本文の合計(選定のみ・評価費用を含まない)。"])
    w3.append([]); w3.append(["独立評価費用(選定費用と別)"]); w3.cell(row=w3.max_row, column=1).font = BOLD
    w3.append(["評価者", "model ID(返却)", "呼出数", "入力token計", "出力token計(推論含む)", "うち推論token", "", "", "評価実費 JPY"])
    for c in w3[w3.max_row]: c.font = BOLD; c.fill = HEAD
    for ev, e in evcost.items():
        w3.append([ev, ", ".join(e["model_returned"]), e["calls"], e["in_tokens"], e["out_tokens"], e["reasoning_tokens"], "", "", round(e["cost_jpy"], 4)])
    sel_total = sum(summary["models"][s]["cost"]["cost_jpy"] for s in C.MODELS)
    ev_total = sum(e["cost_jpy"] for e in evcost.values())
    w3.append([]); w3.append(["Trial実費合計", "", "", "", "", "", "", "", "選定 %.3f + 評価 %.3f = %.3f JPY(台帳paid_calls.jsonl通算 %.3f)" % (sel_total, ev_total, sel_total + ev_total, summary["total_jpy"])])
    w3.append([]); w3.append(["既存Production Key Phrase選定の実費(各runの記録から。モデル/方式が本Trialと異なる[DB hybrid shortlist+canonicalization]ため参考)"]); w3.cell(row=w3.max_row, column=1).font = BOLD
    w3.append(["unit", "選定(selector) JPY", "canonicalization JPY", "区分", "selectorモデル"])
    for c in w3[w3.max_row]: c.font = BOLD; c.fill = HEAD
    tot_sel = 0.0; tot_can = 0.0
    for r in pc:
        w3.append([r["unit_id"], r["selector_jpy"] if r["selector_jpy"] is not None else "未確認", r["canonicalization_jpy"] if r["canonicalization_jpy"] is not None else "未確認", r["kind"], r["selector_model"]])
        tot_sel += r["selector_jpy"] or 0; tot_can += r["canonicalization_jpy"] or 0
    w3.append(["7記事セット合計(14本文)", round(tot_sel, 4), round(tot_can, 4), "実測(selector=Production記録の`cost_jpy`、canonicalization=raw_usage_log token x 単価で再計算。記事Gはgpt-5.6-luna世代の単価で再計算)", ""])
    w3.append(["1記事セット平均(選定+canonicalization)", round((tot_sel + tot_can) / 7, 4)])
    w3.column_dimensions["A"].width = 30; w3.column_dimensions["B"].width = 24; w3.column_dimensions["H"].width = 60
    for col in "CDEFGIJKLMN": w3.column_dimensions[col].width = 16
    # ---- シート4 検証詳細
    w4 = wb.create_sheet("検証詳細")
    w4.append(["選定元", "記事", "レベル", "順位", "候補(出力)", "語数(実測)", "自己申告語数", "自己申告不一致", "5語以内", "本文一致種別", "原文span(置換前)", "文字位置(開始-終了)", "一般化内容(placeholder<-原文)", "重複", "strict合格",
               "主評価 G", "主評価 V", "主評価 E", "主評価 N", "主評価スコア", "参考評価スコア(同系統)", "総合版合格", "備考"])
    for c in w4[1]: c.font = BOLD; c.fill = HEAD; c.alignment = WRAP; c.border = BORDER
    for u in uorder:
        for src in SRC_ORDER:
            for r in [x for x in rows if x["source"] == src and x["unit_id"] == u]:
                per = r["per"]; me = r["main_evs"]
                def mv(k): return avg([per[e][k] for e in me]) if me else None
                uu = next(x for x in units if x["unit_id"] == u)
                w4.append([src_label(src), "記事%s" % uu["article_id"], uu["level"], r["rank"], r["phrase"], r["words"], r["self_count"], "不一致" if r["mismatch"] else "", "OK" if r["within_5"] else "5語超",
                           r["match_type"], r["span"], ("%s-%s" % (r["start"], r["end"])) if r["start"] is not None else "", "; ".join("%s<-%s" % (g["placeholder"], g["original"]) for g in (r["gen"] or [])),
                           "重複(#%s)" % r["dup"] if r["dup"] else "", "合格" if r["strict"] else "不合格", mv("G"), mv("V"), mv("E"), mv("N"), r["main_score"], r["ref_score"], "合格" if r["comp"] else "",
                           r["note"] or ""])
                fl = GREEN if r["strict"] else (YELLOW if r["match_type"] in ("B_GEN_AMBIGUOUS", "NG_INFLECTION_ONLY") and r["within_5"] else RED)
                w4.cell(row=w4.max_row, column=10).fill = fl
    for i, w in enumerate((26, 8, 10, 6, 34, 8, 8, 10, 8, 20, 40, 14, 34, 8, 8, 8, 8, 8, 8, 10, 10, 8, 40), 1): w4.column_dimensions[get_column_letter(i)].width = w
    w4.freeze_panes = "F2"; w4.auto_filter.ref = w4.dimensions
    # ---- シート5 条件・Prompt
    w5 = wb.create_sheet("条件・Prompt")
    prompt = (OUT / "inputs" / "prompt_fixed.txt").read_text(encoding="utf8")
    w5.append(["適用Prompt(全文固定・全6モデル共通・無改変。sha256=%s)" % C.sha_text(prompt)]); w5["A1"].font = BOLD
    w5.append([prompt]); w5.cell(row=2, column=1).alignment = WRAP; w5.row_dimensions[2].height = 330
    w5.append([]); w5.append(["Prompt調整(指示2 事後報告): 変更前=Prompt本文のみ(記事の渡し方は未規定) / 変更後=Prompt本文は1字も変更せず、末尾に『\\n\\n---\\n\\n【記事本文】\\n』+article.md全文を後置 / 理由=記事本文の受け渡し位置が必要(データ受け渡し方式との整合) / 影響なし=Prompt文字列・順序・優先順位・基準は同一(prompt_sha256を全runに記録)、出力形式指示の追加なし、全モデル同一構築。"])
    w5.cell(row=w5.max_row, column=1).alignment = WRAP; w5.row_dimensions[w5.max_row].height = 80
    w5.append([]); w5.append(["各モデルの設定・返却model ID"]); w5.cell(row=w5.max_row, column=1).font = BOLD
    w5.append(["選定元", "指定model ID", "返却model ID", "API", "温度", "推論/thinking", "max出力token", "system", "初回呼出数", "技術的再試行数"])
    for c in w5[w5.max_row]: c.font = BOLD; c.fill = HEAD
    for src, mm in C.MODELS.items():
        rs = [d for (mk, _), d in runs.items() if mk == src]
        ret = sorted({d.get("model_returned") for d in rs if d.get("ok")} - {None})
        retries = sum(len(d["attempts"]) - 1 for d in rs)
        api = {"openai": "OpenAI Responses API", "anthropic": "Anthropic Messages API", "gemini": "Gemini generateContent"}[mm["provider"]]
        w5.append([src_label(src), mm["model_id"], ", ".join(ret), api, "provider既定(未送信)", "provider既定(未送信)。実際の推論token=費用サマリ参照", C.MAX_OUT, "なし", len(rs), retries])
    w5.append([]); w5.append(["評価者(独立評価・モデル名を伏せた7セットを提示)"]); w5.cell(row=w5.max_row, column=1).font = BOLD
    for ev, e in evcost.items(): w5.append([ev, ", ".join(e["model_returned"]), "呼出%d回" % e["calls"]])
    w5.append([]); w5.append(["使用本文の版表"]); w5.cell(row=w5.max_row, column=1).font = BOLD
    w5.append(["unit", "記事", "レベル(内部)", "本文path", "commit", "sha256", "文字数", "語数", "既存KP dir"])
    for c in w5[w5.max_row]: c.font = BOLD; c.fill = HEAD
    for u in units:
        w5.append([u["unit_id"], u["article_name"], "%s(%s)" % (u["level"], u["level_internal"]), u["article_path"], u["article_commit"], u["article_sha256_file"], u["article_chars"], u["article_words"], u["kp_dir"]])
    w5.append([]); w5.append(["単価(USD/1M tokens, input/cached/output[thinking込み])と出典"]); w5.cell(row=w5.max_row, column=1).font = BOLD
    for k, p in C.PRICES.items(): w5.append([k, p[0], p[1], p[2], p[3]])
    w5.column_dimensions["A"].width = 110
    for col in "BCDEFGHI": w5.column_dimensions[col].width = 22
    XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX)
    print("saved", XLSX, "rows", len(rows))
    for src in SRC_ORDER:
        m = summary["models"][src]
        print(src, "n=%d w5=%d conf=%d infl=%d ng=%d strict=%d comp=%d main=%s cost=%.3f" % (m["candidates"], m["within_5"], m["confirmed"], m["infl_only"], m["ng_other"], m["strict"], m["comp"], m["main_avg"], m["cost"]["cost_jpy"]))

if __name__ == "__main__":
    sys.exit(main())
